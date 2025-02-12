import os
import cv2
import sys
import numpy as np
import face_recognition
from datetime import datetime
from openpyxl import Workbook, load_workbook
import tkinter as tk
from tkinter import messagebox

KNOWN_FACES_DIR = "/home/atai/Desktop/project1/known_faces"
known_face_encodings = []
known_face_names = []
marked_today = set()

def load_known_faces():
    global known_face_encodings, known_face_names
    if not os.path.exists(KNOWN_FACES_DIR):
        print(f"Error: The directory '{KNOWN_FACES_DIR}' does not exist. Please create it and add images.")
        return

    for filename in os.listdir(KNOWN_FACES_DIR):
        if filename.endswith('.jpg') or filename.endswith('.png'):
            image_path = os.path.join(KNOWN_FACES_DIR, filename)
            image = face_recognition.load_image_file(image_path)
            face_encoding = face_recognition.face_encodings(image)

            if face_encoding:  # ✅ Ensure encoding is not empty
                known_face_encodings.append(face_encoding[0])
                known_face_names.append(os.path.splitext(filename)[0])

    print("✅ Loaded faces:", known_face_names)

def ensure_sheet_and_headers(sheet_name):
    file_name = 'Attendance.xlsx'
    if not os.path.exists(file_name):
        wb = Workbook()
        wb.create_sheet(sheet_name)
        wb.save(file_name)

    wb = load_workbook(file_name)
    if sheet_name not in wb.sheetnames:
        sheet = wb.create_sheet(sheet_name, 0)
    else:
        sheet = wb[sheet_name]

    if sheet.max_row == 1 or not sheet.cell(1, 1).value:
        headers = ["Name", "Date", "Day", "Time", "Total Periods", "Attended Periods", "Average Attendance"]
        for col_num, header in enumerate(headers, 1):
            sheet.cell(1, col_num, header)

    wb.save(file_name)
    return wb, sheet

def mark_attendance(name, sheet, periods):
    file_name = 'Attendance.xlsx'
    current_time = datetime.now()
    date = current_time.strftime('%d-%m-%y')

    if name in marked_today:
        return

    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[0] == name and row[1] == date:
            return

    next_row = sheet.max_row + 1
    sheet.cell(next_row, 1, name)
    sheet.cell(next_row, 2, date)
    sheet.cell(next_row, 3, current_time.strftime('%A'))
    sheet.cell(next_row, 4, current_time.strftime('%H:%M:%S'))

    update_average_attendance(sheet, name, periods)

    wb = sheet.parent
    wb.save(file_name)
    marked_today.add(name)

    print(f"✅ Attendance marked for {name}")

def update_average_attendance(sheet, name, periods):
    attended_periods = 0
    total_periods = periods

    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[0] == name and row[1] == datetime.now().strftime('%d-%m-%y'):
            attended_periods += total_periods

    average_attendance = (attended_periods / total_periods) * 100 if total_periods else 0

    for row_idx in range(2, sheet.max_row + 1):
        if sheet.cell(row_idx, 1).value == name:
            sheet.cell(row_idx, 5, total_periods)
            sheet.cell(row_idx, 6, attended_periods)
            sheet.cell(row_idx, 7, average_attendance)

def show_alert(name, avg_attendance):
    root = tk.Tk()
    root.withdraw()
    messagebox.showwarning("Low Attendance Alert", f"{name}'s attendance is below 75% ({avg_attendance:.2f}%).")
    root.destroy()

def start_face_recognition(lecture, periods):
    print(f"🎥 Starting face recognition for lecture: {lecture}...")

    wb, sheet = ensure_sheet_and_headers(lecture)
    video_capture = cv2.VideoCapture(0)

    if not video_capture.isOpened():
        print("❌ Error: Could not open camera")
        return

    while True:
        ret, frame = video_capture.read()
        if not ret:
            print("❌ Failed to grab frame. Exiting...")
            break

        rgb_frame = cv2.cvtColor(frame, cv2.COLOR_BGR2RGB)
        face_locations = face_recognition.face_locations(rgb_frame)
        face_encodings = face_recognition.face_encodings(rgb_frame, face_locations)

        for (top, right, bottom, left), face_encoding in zip(face_locations, face_encodings):
            distances = face_recognition.face_distance(known_face_encodings, face_encoding)
            best_match_index = np.argmin(distances) if len(distances) > 0 else None

            name = "Unknown"

            if best_match_index is not None and distances[best_match_index] < 0.5:  
                name = known_face_names[best_match_index]
                mark_attendance(name, sheet, periods)

            avg_attendance = 0
            for row in sheet.iter_rows(min_row=2, values_only=True):
                if row[0] == name:
                    avg_attendance = row[6] if row[6] else 0

            box_color = (0, 255, 0) if avg_attendance >= 75 else (0, 0, 255)
            cv2.rectangle(frame, (left, top), (right, bottom), box_color, 2)

            if avg_attendance < 75:
                show_alert(name, avg_attendance)

            cv2.putText(frame, name, (left, top - 40), cv2.FONT_HERSHEY_SIMPLEX, 0.9, box_color, 2)
            cv2.putText(frame, f"Avg: {avg_attendance:.2f}%", (left, top - 10), cv2.FONT_HERSHEY_SIMPLEX, 0.9, box_color, 2)

        cv2.imshow('Attendance System', frame)

        if cv2.waitKey(1) & 0xFF == ord('q'):
            break

    video_capture.release()
    cv2.destroyAllWindows()

if __name__ == "__main__":
    load_known_faces()  # ✅ Load known faces before starting
    if len(sys.argv) < 3:
        print("Usage: python3 face_recognition_code.py <lecture_name> <periods>")
        sys.exit(1)

    lecture_name = sys.argv[1]
    periods = int(sys.argv[2])

    start_face_recognition(lecture_name, periods)
